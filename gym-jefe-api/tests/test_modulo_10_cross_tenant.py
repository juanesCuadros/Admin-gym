import sys
sys.path.insert(0, ".")
import asyncio
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
import io
import uuid

from httpx import ASGITransport, AsyncClient
import openpyxl
from sqlalchemy import text

from app.core.database import async_session_maker
from app.core.security import create_access_token, hash_password
from app.core.timezone import now_utc, today_local
from app.main import app


async def run_modulo_10_full_test_suite():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        print("=== INICIANDO SUITE COMPLETA MÓDULO 10: REPORTES Y MÉTRICAS (RF-41 a RF-44) ===")

        # 1. SETUP TENANTS (Gym A y Gym B)
        prefix = uuid.uuid4().hex[:6]
        gym_a_id = uuid.uuid4()
        gym_b_id = uuid.uuid4()
        staff_jefe_a_id = uuid.uuid4()
        staff_recep_a_id = uuid.uuid4()
        staff_coach_a_id = uuid.uuid4()
        staff_jefe_b_id = uuid.uuid4()
        staff_recep_b_id = uuid.uuid4()

        pwd_hash = hash_password("Password123!")
        hoy = today_local()

        dep_a1_id = uuid.uuid4()  # Deportista con membresía por vencer
        dep_a2_id = uuid.uuid4()  # Deportista inactivo (sin checkin reciente)
        dep_a3_id = uuid.uuid4()  # Deportista constante (con checkin hoy)
        plan_a_id = uuid.uuid4()
        mem_a1_id = uuid.uuid4()
        mem_a2_id = uuid.uuid4()
        mem_a3_id = uuid.uuid4()
        turno_a_id = uuid.uuid4()
        prod_a_id = uuid.uuid4()

        async with async_session_maker() as session:
            # Configurar Gym A
            await session.execute(text("SELECT set_config('app.gimnasio_id', :gym_id, true)"), {"gym_id": str(gym_a_id)})
            await session.execute(text("""
                INSERT INTO platform.tenant (id, nombre, subdominio, activo, dias_gracia_mora, dias_umbral_por_vencer)
                VALUES (:id_a, :nom_a, :sub_a, true, 3, 5)
            """), {"id_a": gym_a_id, "nom_a": f"Gym Rep A {prefix}", "sub_a": f"rep-a-{prefix}"})

            # Staff Gym A
            await session.execute(text("""
                INSERT INTO platform.staff (id, gimnasio_id, rol, correo, hash_password, nombre, activo)
                VALUES (:id, :gym, 'jefe', :mail, :pwd, 'Jefe Rep A', true)
            """), {"id": staff_jefe_a_id, "gym": gym_a_id, "mail": f"jefe_a_{prefix}@gymos.co", "pwd": pwd_hash})

            await session.execute(text("""
                INSERT INTO platform.staff (id, gimnasio_id, rol, correo, hash_password, nombre, activo)
                VALUES (:id, :gym, 'recepcionista', :mail, :pwd, 'Recep Rep A', true)
            """), {"id": staff_recep_a_id, "gym": gym_a_id, "mail": f"recep_a_{prefix}@gymos.co", "pwd": pwd_hash})

            await session.execute(text("""
                INSERT INTO platform.staff (id, gimnasio_id, rol, correo, hash_password, nombre, activo)
                VALUES (:id, :gym, 'entrenador', :mail, :pwd, 'Coach Rep A', true)
            """), {"id": staff_coach_a_id, "gym": gym_a_id, "mail": f"coach_a_{prefix}@gymos.co", "pwd": pwd_hash})

            from app.modules.auth.service import AuthService
            await AuthService._seed_permisos_en_transaccion(session, gym_a_id)

            # Deportistas en Gym A
            await session.execute(text("""
                INSERT INTO platform.deportistas (id, gimnasio_id, nombre, documento, correo, telefono, activo)
                VALUES 
                (:d1, :gym, 'Atleta Por Vencer', 'DOC-VENCER', 'vencer@gymos.co', '3001112233', true),
                (:d2, :gym, 'Atleta Inactivo', 'DOC-INACTIVO', 'inactivo@gymos.co', '3004445566', true),
                (:d3, :gym, 'Atleta Frecuente', 'DOC-FRECUENTE', 'frecuente@gymos.co', '3007778899', true)
            """), {"gym": gym_a_id, "d1": dep_a1_id, "d2": dep_a2_id, "d3": dep_a3_id})

            # Plan en Gym A
            await session.execute(text("""
                INSERT INTO platform.planes (id, gimnasio_id, nombre, precio, duracion_dias, tipo, activo)
                VALUES (:p_id, :gym, 'Plan Mensual VIP', 120000.00, 30, 'individual', true)
            """), {"p_id": plan_a_id, "gym": gym_a_id})

            # Membresías en Gym A
            # 1. Por vencer en 3 días (umbral tenant es 5)
            await session.execute(text("""
                INSERT INTO platform.membresias (id, gimnasio_id, deportista_id, plan_id, fecha_inicio, fecha_vencimiento, cancelada)
                VALUES (:m1, :gym, :d1, :p_id, :ini, :venc, false)
            """), {"m1": mem_a1_id, "gym": gym_a_id, "d1": dep_a1_id, "p_id": plan_a_id, "ini": hoy - timedelta(days=27), "venc": hoy + timedelta(days=3)})

            # 2. Inactivo: inició hace 25 días, vence en 15 días, último check-in hace 20 días
            await session.execute(text("""
                INSERT INTO platform.membresias (id, gimnasio_id, deportista_id, plan_id, fecha_inicio, fecha_vencimiento, cancelada)
                VALUES (:m2, :gym, :d2, :p_id, :ini, :venc, false)
            """), {"m2": mem_a2_id, "gym": gym_a_id, "d2": dep_a2_id, "p_id": plan_a_id, "ini": hoy - timedelta(days=25), "venc": hoy + timedelta(days=15)})

            # 3. Frecuente: vence en 20 días
            await session.execute(text("""
                INSERT INTO platform.membresias (id, gimnasio_id, deportista_id, plan_id, fecha_inicio, fecha_vencimiento, cancelada)
                VALUES (:m3, :gym, :d3, :p_id, :ini, :venc, false)
            """), {"m3": mem_a3_id, "gym": gym_a_id, "d3": dep_a3_id, "p_id": plan_a_id, "ini": hoy - timedelta(days=10), "venc": hoy + timedelta(days=20)})

            # Turno de caja en Gym A
            await session.execute(text("""
                INSERT INTO platform.turnos_caja (id, gimnasio_id, staff_id, base_inicial, estado, abierto_en)
                VALUES (:t_id, :gym, :s_id, 50000.00, 'abierto', now())
            """), {"t_id": turno_a_id, "gym": gym_a_id, "s_id": staff_recep_a_id})

            # Producto en Gym A
            await session.execute(text("""
                INSERT INTO platform.productos (id, gimnasio_id, nombre, precio, stock, activo)
                VALUES (:pr_id, :gym, 'Bebida Hidratante', 5000.00, 50, true)
            """), {"pr_id": prod_a_id, "gym": gym_a_id})

            # Ventas en Gym A
            # Venta 1 activa: $50,000 en efectivo
            v1_id = uuid.uuid4()
            await session.execute(text("""
                INSERT INTO platform.ventas (id, gimnasio_id, turno_id, total, metodo, tipo_medio, anulada, idempotency_key, registrada_por, created_at)
                VALUES (:v1, :gym, :t_id, 50000.00, 'efectivo', 'efectivo', false, :idemp1, :sid, now())
            """), {"v1": v1_id, "gym": gym_a_id, "t_id": turno_a_id, "idemp1": f"v1_{prefix}", "sid": staff_recep_a_id})

            await session.execute(text("""
                INSERT INTO platform.venta_items (gimnasio_id, venta_id, tipo, producto_id, descripcion, cantidad, precio_unitario, subtotal)
                VALUES (:gym, :v1, 'producto', :pr_id, 'Bebida Hidratante', 10, 5000.00, 50000.00)
            """), {"gym": gym_a_id, "v1": v1_id, "pr_id": prod_a_id})

            # Venta 2 ANULADA: $25,000 con tarjeta
            v2_id = uuid.uuid4()
            await session.execute(text("""
                INSERT INTO platform.ventas (id, gimnasio_id, turno_id, total, metodo, tipo_medio, anulada, motivo_anulacion, anulada_en, anulada_por, idempotency_key, registrada_por, created_at)
                VALUES (:v2, :gym, :t_id, 25000.00, 'tarjeta', 'otro', true, 'Error de digitación en datafono', now(), :sid, :idemp2, :sid, now())
            """), {"v2": v2_id, "gym": gym_a_id, "t_id": turno_a_id, "sid": staff_recep_a_id, "idemp2": f"v2_{prefix}"})

            # Pago de membresía: $120,000 por Nequi
            pm_id = uuid.uuid4()
            await session.execute(text("""
                INSERT INTO platform.pagos_membresia (id, gimnasio_id, turno_id, membresia_id, monto, metodo, tipo_medio, dias_agregados, anulado, idempotency_key, registrado_por, created_at)
                VALUES (:pm, :gym, :t_id, :m3, 120000.00, 'nequi', 'otro', 30, false, :idemp_pm, :sid, now())
            """), {"pm": pm_id, "gym": gym_a_id, "t_id": turno_a_id, "m3": mem_a3_id, "idemp_pm": f"pm_{prefix}", "sid": staff_recep_a_id})

            # Devolución: $10,000 en efectivo
            dev_id = uuid.uuid4()
            await session.execute(text("""
                INSERT INTO platform.devoluciones (id, gimnasio_id, venta_original_id, turno_venta_original_id, turno_devolucion_id, monto, tipo_medio_reembolso, motivo, registrada_por, created_at)
                VALUES (:dev, :gym, :v1, :t_id, :t_id, 10000.00, 'efectivo', 'Producto defectuoso', :sid, now())
            """), {"dev": dev_id, "gym": gym_a_id, "v1": v1_id, "t_id": turno_a_id, "sid": staff_recep_a_id})

            # Check-ins en Gym A
            # 1. Atleta Frecuente (hoy)
            await session.execute(text("""
                INSERT INTO platform.checkins (gimnasio_id, deportista_id, tipo, metodo, resultado, registrado_por, ts_utc)
                VALUES (:gym, :d3, 'ingreso', 'manual', 'abrio', :sid, now())
            """), {"gym": gym_a_id, "d3": dep_a3_id, "sid": staff_recep_a_id})

            # 2. Atleta Inactivo (hace 20 días)
            await session.execute(text("""
                INSERT INTO platform.checkins (gimnasio_id, deportista_id, tipo, metodo, resultado, registrado_por, ts_utc)
                VALUES (:gym, :d2, 'ingreso', 'manual', 'abrio', :sid, now() - interval '20 days')
            """), {"gym": gym_a_id, "d2": dep_a2_id, "sid": staff_recep_a_id})

            # 3. Checkin con alerta mora
            await session.execute(text("""
                INSERT INTO platform.checkins (gimnasio_id, deportista_id, tipo, metodo, resultado, registrado_por, ts_utc)
                VALUES (:gym, :d1, 'ingreso', 'manual', 'alerta_mora', :sid, now())
            """), {"gym": gym_a_id, "d1": dep_a1_id, "sid": staff_recep_a_id})

            # 4. Checkin denegado
            await session.execute(text("""
                INSERT INTO platform.checkins (gimnasio_id, deportista_id, tipo, metodo, resultado, registrado_por, ts_utc)
                VALUES (:gym, :d1, 'ingreso', 'manual', 'negado', :sid, now())
            """), {"gym": gym_a_id, "d1": dep_a1_id, "sid": staff_recep_a_id})

            # 5. Cortesía
            await session.execute(text("""
                INSERT INTO platform.checkins (gimnasio_id, deportista_id, tipo, metodo, resultado, motivo_cortesia, registrado_por, ts_utc)
                VALUES (:gym, null, 'cortesia', 'manual', 'abrio', 'Cortesía de bienvenida familiar', :sid, now())
            """), {"gym": gym_a_id, "sid": staff_recep_a_id})

            # Configurar Tenant B y Staff Gym B (Completamente vacío de transacciones y deportistas)
            await session.execute(text("SELECT set_config('app.gimnasio_id', :gym_id, true)"), {"gym_id": str(gym_b_id)})
            await session.execute(text("""
                INSERT INTO platform.tenant (id, nombre, subdominio, activo, dias_gracia_mora, dias_umbral_por_vencer)
                VALUES (:id_b, :nom_b, :sub_b, true, 3, 5)
            """), {"id_b": gym_b_id, "nom_b": f"Gym Rep B {prefix}", "sub_b": f"rep-b-{prefix}"})

            await session.execute(text("""
                INSERT INTO platform.staff (id, gimnasio_id, rol, correo, hash_password, nombre, activo)
                VALUES (:id, :gym, 'jefe', :mail, :pwd, 'Jefe Rep B', true)
            """), {"id": staff_jefe_b_id, "gym": gym_b_id, "mail": f"jefe_b_{prefix}@gymos.co", "pwd": pwd_hash})

            await session.execute(text("""
                INSERT INTO platform.staff (id, gimnasio_id, rol, correo, hash_password, nombre, activo)
                VALUES (:id, :gym, 'recepcionista', :mail, :pwd, 'Recep Rep B', true)
            """), {"id": staff_recep_b_id, "gym": gym_b_id, "mail": f"recep_b_{prefix}@gymos.co", "pwd": pwd_hash})

            await AuthService._seed_permisos_en_transaccion(session, gym_b_id)
            await session.commit()

        # JWT Tokens
        token_jefe_a = create_access_token(subject=str(staff_jefe_a_id), gym_id=str(gym_a_id), role="jefe")
        token_coach_a = create_access_token(subject=str(staff_coach_a_id), gym_id=str(gym_a_id), role="entrenador")
        token_recep_a = create_access_token(subject=str(staff_recep_a_id), gym_id=str(gym_a_id), role="recepcionista")
        token_jefe_b = create_access_token(subject=str(staff_jefe_b_id), gym_id=str(gym_b_id), role="jefe")
        token_recep_b = create_access_token(subject=str(staff_recep_b_id), gym_id=str(gym_b_id), role="recepcionista")

        headers_jefe_a = {"Authorization": f"Bearer {token_jefe_a}"}
        headers_coach_a = {"Authorization": f"Bearer {token_coach_a}"}
        headers_recep_a = {"Authorization": f"Bearer {token_recep_a}"}
        headers_jefe_b = {"Authorization": f"Bearer {token_jefe_b}"}

        print("--- [1/7] TEST RF-41: REPORTE DE INGRESOS, EXCLUSIÓN DE ANULADAS Y TRAZABILIDAD ---")
        # Jefe A consulta ingresos del mes
        res_ing_a = await client.get("/api/v1/reportes/ingresos", headers=headers_jefe_a)
        assert res_ing_a.status_code == 200, f"Error consultando ingresos: {res_ing_a.text}"
        data_ing_a = res_ing_a.json()

        # Validaciones matemáticas exactas:
        # Bruto = $50,000 (venta 1) + $120,000 (pago membresía) = $170,000.00
        assert Decimal(str(data_ing_a["total_ingresos_bruto"])) == Decimal("170000.00")
        # Devoluciones = $10,000.00
        assert Decimal(str(data_ing_a["total_devoluciones"])) == Decimal("10000.00")
        # Neto = $170,000 - $10,000 = $160,000.00
        assert Decimal(str(data_ing_a["total_ingresos_neto"])) == Decimal("160000.00")
        # Anuladas = $25,000.00 (excluido de neto y bruto, pero reportado)
        assert Decimal(str(data_ing_a["total_anulado"])) == Decimal("25000.00")

        # Desglose por método
        assert Decimal(str(data_ing_a["desglose_por_metodo"]["efectivo"])) == Decimal("50000.00")
        assert Decimal(str(data_ing_a["desglose_por_metodo"]["nequi"])) == Decimal("120000.00")

        # Desglose por fuente
        assert Decimal(str(data_ing_a["desglose_por_fuente"]["membresias"])) == Decimal("120000.00")
        assert Decimal(str(data_ing_a["desglose_por_fuente"]["productos"])) == Decimal("50000.00")

        # Verificar lista de transacciones con la fila anulada
        txs = data_ing_a["transacciones"]
        tx_anulada = next(t for t in txs if t["id"] == str(v2_id))
        assert tx_anulada["anulado"] is True
        assert tx_anulada["motivo_anulacion"] == "Error de digitación en datafono"

        print("--- [2/7] TEST RF-41: AISLAMIENTO CROSS-TENANT EN AGREGACIONES DE INGRESOS ---")
        # Jefe B consulta ingresos: debe recibir EXACTAMENTE $0.00 en todos los acumuladores y 0 transacciones
        res_ing_b = await client.get("/api/v1/reportes/ingresos", headers=headers_jefe_b)
        assert res_ing_b.status_code == 200
        data_ing_b = res_ing_b.json()
        assert Decimal(str(data_ing_b["total_ingresos_bruto"])) == Decimal("0.00")
        assert Decimal(str(data_ing_b["total_ingresos_neto"])) == Decimal("0.00")
        assert Decimal(str(data_ing_b["total_anulado"])) == Decimal("0.00")
        assert len(data_ing_b["transacciones"]) == 0

        print("--- [3/7] TEST RF-42: MEMBRESÍAS POR VENCER CON DOMINIO CENTRALIZADO ---")
        # Jefe A consulta membresías por vencer (umbral tenant = 5 días)
        res_venc_a = await client.get("/api/v1/reportes/membresias-por-vencer", headers=headers_jefe_a)
        assert res_venc_a.status_code == 200
        data_venc_a = res_venc_a.json()
        assert data_venc_a["umbral_dias_aplicado"] == 5
        assert data_venc_a["total"] >= 1
        item_venc = next(it for it in data_venc_a["items"] if it["documento"] == "DOC-VENCER")
        assert item_venc["dias_restantes"] == 3
        assert item_venc["plan_nombre"] == "Plan Mensual VIP"
        assert item_venc["telefono"] == "3001112233"

        # Cross-Tenant: Gym B consulta membresías por vencer -> 0 registros
        res_venc_b = await client.get("/api/v1/reportes/membresias-por-vencer", headers=headers_jefe_b)
        assert res_venc_b.status_code == 200
        assert res_venc_b.json()["total"] == 0
        assert len(res_venc_b.json()["items"]) == 0

        print("--- [4/7] TEST RF-43: ASISTENCIA Y AFLUENCIA POR HORARIOS ---")
        # Jefe A consulta asistencia
        res_asist_a = await client.get(f"/api/v1/reportes/asistencia?fecha_inicio={hoy - timedelta(days=25)}&fecha_fin={hoy}", headers=headers_jefe_a)
        assert res_asist_a.status_code == 200
        data_asist_a = res_asist_a.json()
        assert data_asist_a["total_accesos"] == 5
        assert data_asist_a["accesos_permitidos"] == 4
        assert data_asist_a["accesos_denegados"] == 1
        assert data_asist_a["cortesias"] == 1
        assert len(data_asist_a["afluencia_por_hora"]) == 24
        assert len(data_asist_a["afluencia_por_dia_semana"]) == 7

        # Cross-Tenant: Gym B consulta asistencia -> Cero accesos en todos los campos
        res_asist_b = await client.get("/api/v1/reportes/asistencia", headers=headers_jefe_b)
        assert res_asist_b.status_code == 200
        data_asist_b = res_asist_b.json()
        assert data_asist_b["total_accesos"] == 0
        assert data_asist_b["accesos_permitidos"] == 0
        assert data_asist_b["accesos_denegados"] == 0

        print("--- [5/7] TEST RF-43 EXTENSIÓN: DEPORTISTAS INACTIVOS (RETENCIÓN) ---")
        # Jefe A consulta deportistas inactivos (umbral 15 días)
        res_inac_a = await client.get("/api/v1/reportes/deportistas-inactivos?dias_sin_asistencia=15", headers=headers_jefe_a)
        assert res_inac_a.status_code == 200
        data_inac_a = res_inac_a.json()
        assert data_inac_a["total"] >= 1
        item_inac = next(it for it in data_inac_a["items"] if it["documento"] == "DOC-INACTIVO")
        assert item_inac["dias_sin_asistir"] >= 15
        assert item_inac["telefono"] == "3004445566"
        assert item_inac["plan_nombre"] == "Plan Mensual VIP"

        # Cross-Tenant: Gym B no tiene deportistas inactivos
        res_inac_b = await client.get("/api/v1/reportes/deportistas-inactivos", headers=headers_jefe_b)
        assert res_inac_b.status_code == 200
        assert res_inac_b.json()["total"] == 0

        print("--- [6/7] TEST RF-44: EXPORTACIÓN EN MEMORIA A CSV Y EXCEL (OPENPYXL) ---")
        # 1. Exportar Ingresos a CSV con UTF-8 BOM
        res_csv = await client.get("/api/v1/reportes/exportar?tipo_reporte=ingresos&formato=csv", headers=headers_jefe_a)
        assert res_csv.status_code == 200
        assert "text/csv" in res_csv.headers["content-type"]
        assert "attachment; filename=" in res_csv.headers["content-disposition"]
        # Validar BOM UTF-8 (\ufeff)
        csv_text = res_csv.content.decode("utf-8")
        assert csv_text.startswith("\ufeff")
        assert "50000.0" in csv_text or "50000" in csv_text
        assert "Error de digitación en datafono" in csv_text

        # 2. Exportar Membresías por Vencer a Excel (.xlsx) y validar con openpyxl
        res_xlsx = await client.get("/api/v1/reportes/exportar?tipo_reporte=membresias_por_vencer&formato=excel", headers=headers_jefe_a)
        assert res_xlsx.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res_xlsx.headers["content-type"]
        wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
        ws = wb.active
        assert ws.title == "Membresías Por Vencer"
        # Comprobar cabecera y datos
        rows_excel = list(ws.iter_rows(values_only=True))
        assert "Nombre Deportista" in rows_excel[0]
        assert any(r[0] == "Atleta Por Vencer" for r in rows_excel[1:])

        # 3. Exportar Deportistas Inactivos a Excel
        res_xlsx_inac = await client.get("/api/v1/reportes/exportar?tipo_reporte=deportistas_inactivos&formato=excel", headers=headers_jefe_a)
        assert res_xlsx_inac.status_code == 200
        wb_inac = openpyxl.load_workbook(io.BytesIO(res_xlsx_inac.content))
        rows_inac = list(wb_inac.active.iter_rows(values_only=True))
        assert any(r[0] == "Atleta Inactivo" for r in rows_inac[1:])

        # 4. Cross-Tenant: Gym B exporta a CSV -> Archivo contiene cabecera pero NINGUNA fila de datos de Gym A
        res_csv_b = await client.get("/api/v1/reportes/exportar?tipo_reporte=ingresos&formato=csv", headers=headers_jefe_b)
        assert res_csv_b.status_code == 200
        csv_b_lines = [l for l in res_csv_b.content.decode("utf-8").strip().splitlines() if l]
        assert len(csv_b_lines) == 1, f"Gym B no debería tener filas de datos en exportación, encontró: {csv_b_lines}"

        print("--- [7/7] TEST RECHAZO RBAC EN REPORTES PARA ROLES SUBORDINADOS ---")
        # Recepcionista intenta consultar ingresos -> 403 PERMISO_DENEGADO
        res_recep_forbidden = await client.get("/api/v1/reportes/ingresos", headers=headers_recep_a)
        assert res_recep_forbidden.status_code == 403
        assert res_recep_forbidden.json()["error"]["codigo"] == "PERMISO_DENEGADO"

        # Entrenador intenta consultar asistencia -> 403 PERMISO_DENEGADO
        res_coach_forbidden = await client.get("/api/v1/reportes/asistencia", headers=headers_coach_a)
        assert res_coach_forbidden.status_code == 403
        assert res_coach_forbidden.json()["error"]["codigo"] == "PERMISO_DENEGADO"

        # Recepcionista intenta exportar reporte -> 403 PERMISO_DENEGADO
        res_exp_forbidden = await client.get("/api/v1/reportes/exportar?tipo_reporte=ingresos&formato=csv", headers=headers_recep_a)
        assert res_exp_forbidden.status_code == 403
        assert res_exp_forbidden.json()["error"]["codigo"] == "PERMISO_DENEGADO"

        print("=== TODAS LAS PRUEBAS DEL MÓDULO 10 (REPORTES Y MÉTRICAS) PASARON EXITOSAMENTE ===")


if __name__ == "__main__":
    asyncio.run(run_modulo_10_full_test_suite())
