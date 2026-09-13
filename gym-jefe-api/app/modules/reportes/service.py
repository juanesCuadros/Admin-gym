"""Lógica de negocio y agregaciones para el módulo de Reportes y Métricas (RF-41 a RF-44)"""
import csv
import io
import math
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from typing import Any, Dict, List, Optional
from uuid import UUID

from fastapi import HTTPException, Response, status
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.timezone import get_local_day_range_utc, now_local, now_utc, today_local
from app.modules.membresias.domain import MembresiaDomainService
from app.modules.reportes.schemas import (
    AfluenciaDiaSemanaDto,
    AfluenciaHoraDto,
    DeportistaInactivoItemDto,
    MembresiaPorVencerItemDto,
    ReporteAsistenciaResponse,
    ReporteDeportistasInactivosResponse,
    ReporteIngresosResponse,
    ReporteMembresiasPorVencerResponse,
    TopDeportistaAsistenciaDto,
    TransaccionReporteDto,
)


class ReportesService:
    """Servicio de generación, agregación y exportación de reportes operativos y financieros."""

    # =========================================================================
    # 1. RF-41: REPORTE FINANCIERO DE INGRESOS
    # =========================================================================

    @classmethod
    async def obtener_reporte_ingresos(
        cls,
        session: AsyncSession,
        gym_id: UUID,
        fecha_inicio: Optional[date] = None,
        fecha_fin: Optional[date] = None,
        metodo: Optional[str] = None,
        fuente: Optional[str] = None,
    ) -> ReporteIngresosResponse:
        hoy = today_local()
        if fecha_fin is None:
            fecha_fin = hoy
        if fecha_inicio is None:
            fecha_inicio = fecha_fin.replace(day=1)

        if fecha_inicio > fecha_fin:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={"codigo": "RANGO_FECHAS_INVALIDO", "mensaje": "La fecha de inicio no puede ser posterior a la fecha de fin"}
            )

        # Corte exacto en UTC alineado a medianoche de America/Bogota
        inicio_utc = get_local_day_range_utc(fecha_inicio)[0]
        fin_utc = get_local_day_range_utc(fecha_fin)[1]

        # 1. Consultar ventas de mostrador (productos, pases)
        q_ventas = text("""
            SELECT v.id, v.total, v.metodo, v.tipo_medio, v.anulada, v.motivo_anulacion, v.anulada_en, v.created_at,
                   s.nombre as staff_nombre,
                   COALESCE((
                       SELECT jsonb_agg(jsonb_build_object(
                           'tipo', vi.tipo, 'descripcion', vi.descripcion, 'cantidad', vi.cantidad, 'subtotal', vi.subtotal
                       ))
                       FROM platform.venta_items vi
                       WHERE vi.venta_id = v.id
                   ), '[]'::jsonb) as items
            FROM platform.ventas v
            LEFT JOIN platform.staff s ON s.id = v.registrada_por
            WHERE v.gimnasio_id = :gym_id AND v.created_at >= :inicio_utc AND v.created_at <= :fin_utc
            ORDER BY v.created_at DESC
        """)
        res_ventas = await session.execute(q_ventas, {"gym_id": gym_id, "inicio_utc": inicio_utc, "fin_utc": fin_utc})
        ventas_rows = res_ventas.mappings().all()

        # 2. Consultar pagos de membresía
        q_pagos = text("""
            SELECT pm.id, pm.monto, pm.metodo, pm.tipo_medio, pm.anulado, pm.motivo_anulacion, pm.anulado_en, pm.created_at,
                   s.nombre as staff_nombre, p.nombre as plan_nombre, d.nombre as deportista_nombre
            FROM platform.pagos_membresia pm
            LEFT JOIN platform.staff s ON s.id = pm.registrado_por
            JOIN platform.membresias m ON m.id = pm.membresia_id
            JOIN platform.planes p ON p.id = m.plan_id
            JOIN platform.deportistas d ON d.id = m.deportista_id
            WHERE pm.gimnasio_id = :gym_id AND pm.created_at >= :inicio_utc AND pm.created_at <= :fin_utc
            ORDER BY pm.created_at DESC
        """)
        res_pagos = await session.execute(q_pagos, {"gym_id": gym_id, "inicio_utc": inicio_utc, "fin_utc": fin_utc})
        pagos_rows = res_pagos.mappings().all()

        # 3. Consultar devoluciones
        q_devs = text("""
            SELECT d.id, d.monto, d.tipo_medio_reembolso, d.motivo, d.created_at,
                   s.nombre as staff_nombre
            FROM platform.devoluciones d
            LEFT JOIN platform.staff s ON s.id = d.registrada_por
            WHERE d.gimnasio_id = :gym_id AND d.created_at >= :inicio_utc AND d.created_at <= :fin_utc
            ORDER BY d.created_at DESC
        """)
        res_devs = await session.execute(q_devs, {"gym_id": gym_id, "inicio_utc": inicio_utc, "fin_utc": fin_utc})
        devs_rows = res_devs.mappings().all()

        # Acumuladores
        total_ingresos_bruto = Decimal("0.00")
        total_anulado = Decimal("0.00")
        total_devoluciones = Decimal("0.00")
        desglose_por_metodo: Dict[str, Decimal] = {}
        desglose_por_fuente: Dict[str, Decimal] = {
            "membresias": Decimal("0.00"),
            "productos": Decimal("0.00"),
            "otros_mostrador": Decimal("0.00"),
        }

        transacciones: List[TransaccionReporteDto] = []

        # Procesar Ventas
        for v in ventas_rows:
            total_v = Decimal(str(v["total"]))
            metodo_v = str(v["metodo"]).lower().strip()
            es_anulada = bool(v["anulada"])

            # Desglose de ítems para detalle
            items_list = v["items"] or []
            desc_items = ", ".join([f"{it['cantidad']}x {it['descripcion']}" for it in items_list]) or "Venta mostrador"

            if es_anulada:
                total_anulado += total_v
            else:
                total_ingresos_bruto += total_v
                desglose_por_metodo[metodo_v] = desglose_por_metodo.get(metodo_v, Decimal("0.00")) + total_v

                # Categorizar ítems por fuente
                for it in items_list:
                    sub_it = Decimal(str(it["subtotal"]))
                    if it["tipo"] == "producto":
                        desglose_por_fuente["productos"] += sub_it
                    else:
                        desglose_por_fuente["otros_mostrador"] += sub_it

            transacciones.append(TransaccionReporteDto(
                id=v["id"],
                fecha_hora=v["created_at"],
                tipo="venta_mostrador",
                descripcion=desc_items,
                monto=total_v,
                metodo=metodo_v,
                tipo_medio=v["tipo_medio"],
                anulado=es_anulada,
                motivo_anulacion=v["motivo_anulacion"],
                registrado_por_nombre=v["staff_nombre"]
            ))

        # Procesar Pagos de Membresía
        for p in pagos_rows:
            monto_p = Decimal(str(p["monto"]))
            metodo_p = str(p["metodo"]).lower().strip()
            es_anulado = bool(p["anulado"])
            desc_pago = f"Pago membresía: {p['plan_nombre']} ({p['deportista_nombre']})"

            if es_anulado:
                total_anulado += monto_p
            else:
                total_ingresos_bruto += monto_p
                desglose_por_metodo[metodo_p] = desglose_por_metodo.get(metodo_p, Decimal("0.00")) + monto_p
                desglose_por_fuente["membresias"] += monto_p

            transacciones.append(TransaccionReporteDto(
                id=p["id"],
                fecha_hora=p["created_at"],
                tipo="pago_membresia",
                descripcion=desc_pago,
                monto=monto_p,
                metodo=metodo_p,
                tipo_medio=p["tipo_medio"],
                anulado=es_anulado,
                motivo_anulacion=p["motivo_anulacion"],
                registrado_por_nombre=p["staff_nombre"]
            ))

        # Procesar Devoluciones
        for d in devs_rows:
            monto_d = Decimal(str(d["monto"]))
            total_devoluciones += monto_d
            transacciones.append(TransaccionReporteDto(
                id=d["id"],
                fecha_hora=d["created_at"],
                tipo="devolucion",
                descripcion=f"Devolución: {d['motivo']}",
                monto=-monto_d,
                metodo=f"devolucion_{d['tipo_medio_reembolso']}",
                tipo_medio=d["tipo_medio_reembolso"],
                anulado=False,
                motivo_anulacion=None,
                registrado_por_nombre=d["staff_nombre"]
            ))

        total_ingresos_neto = total_ingresos_bruto - total_devoluciones

        # Filtrar transacciones en memoria si se especificó método o fuente
        if metodo:
            filtro_m = metodo.lower().strip()
            transacciones = [t for t in transacciones if t.metodo == filtro_m]

        if fuente:
            filtro_f = fuente.lower().strip()
            if filtro_f == "membresias":
                transacciones = [t for t in transacciones if t.tipo == "pago_membresia"]
            elif filtro_f == "mostrador":
                transacciones = [t for t in transacciones if t.tipo == "venta_mostrador"]

        # Ordenar transacciones por fecha desc
        transacciones.sort(key=lambda t: t.fecha_hora, reverse=True)

        return ReporteIngresosResponse(
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_fin,
            total_ingresos_bruto=total_ingresos_bruto,
            total_devoluciones=total_devoluciones,
            total_ingresos_neto=total_ingresos_neto,
            total_anulado=total_anulado,
            desglose_por_metodo=desglose_por_metodo,
            desglose_por_fuente=desglose_por_fuente,
            transacciones=transacciones
        )

    # =========================================================================
    # 2. RF-42: REPORTE DE MEMBRESÍAS POR VENCER (ACCIONABLE)
    # =========================================================================

    @classmethod
    async def obtener_membresias_por_vencer(
        cls,
        session: AsyncSession,
        gym_id: UUID,
        dias: Optional[int] = None,
        pagina: int = 1,
        limite: int = 50,
    ) -> ReporteMembresiasPorVencerResponse:
        hoy = today_local()
        pagina = max(1, pagina)
        limite = max(1, min(limite, 10000))

        # 1. Obtener configuración del tenant
        q_conf = text("SELECT dias_gracia_mora, dias_umbral_por_vencer FROM platform.tenant WHERE id = :gym_id")
        res_conf = await session.execute(q_conf, {"gym_id": gym_id})
        row_conf = res_conf.mappings().first()
        gracia = row_conf["dias_gracia_mora"] if row_conf else 3
        umbral_tenant = row_conf["dias_umbral_por_vencer"] if row_conf else 5

        umbral_efectivo = dias if dias is not None else umbral_tenant

        # 2. Consultar membresías vigentes de deportistas activos
        q_membresias = text("""
            SELECT m.id, m.deportista_id, d.nombre as deportista_nombre, d.documento as deportista_documento,
                   d.correo as deportista_correo, d.telefono as deportista_telefono, d.activo as deportista_activo,
                   m.plan_id, p.nombre as plan_nombre, m.fecha_inicio, m.fecha_vencimiento, m.cancelada
            FROM platform.membresias m
            JOIN platform.deportistas d ON d.id = m.deportista_id
            JOIN platform.planes p ON p.id = m.plan_id
            WHERE m.gimnasio_id = :gym_id
              AND d.deleted_at IS NULL
              AND d.activo = true
              AND m.cancelada = false
              AND m.fecha_vencimiento >= :hoy
            ORDER BY m.fecha_vencimiento ASC
        """)
        res_m = await session.execute(q_membresias, {"gym_id": gym_id, "hoy": hoy})
        raw_items = [dict(r) for r in res_m.mappings().all()]

        # 3. Enriquecer con el Servicio de Dominio Único
        enriquecidos = await MembresiaDomainService.enriquecer_membresias_batch(
            session=session,
            gym_id=gym_id,
            items=raw_items,
            dias_gracia_mora=gracia,
            dias_umbral_por_vencer=umbral_efectivo
        )

        # 4. Filtrar los que están en umbral de "por vencer"
        items_filtrados = [
            item for item in enriquecidos
            if item.get("estado_calculado") == "por_vencer" or (0 <= item.get("dias_restantes_o_vencido", 999) <= umbral_efectivo)
        ]

        total = len(items_filtrados)
        total_paginas = max(1, math.ceil(total / limite))
        offset = (pagina - 1) * limite
        pagina_items = items_filtrados[offset : offset + limite]

        dtos = [
            MembresiaPorVencerItemDto(
                membresia_id=it["id"],
                deportista_id=it["deportista_id"],
                nombre=it["deportista_nombre"],
                documento=it["deportista_documento"],
                correo=it.get("deportista_correo"),
                telefono=it.get("deportista_telefono"),
                plan_id=it["plan_id"],
                plan_nombre=it["plan_nombre"],
                fecha_inicio=it["fecha_inicio"],
                fecha_vencimiento=it["fecha_vencimiento"],
                dias_restantes=it["dias_restantes_o_vencido"],
                congelamiento_activo=bool(it.get("congelamiento_activo", False))
            )
            for it in pagina_items
        ]

        return ReporteMembresiasPorVencerResponse(
            umbral_dias_aplicado=umbral_efectivo,
            total=total,
            pagina=pagina,
            limite=limite,
            total_paginas=total_paginas,
            items=dtos
        )

    # =========================================================================
    # 3. RF-43: REPORTE DE ASISTENCIA Y AFLUENCIA POR HORARIOS
    # =========================================================================

    @classmethod
    async def obtener_reporte_asistencia(
        cls,
        session: AsyncSession,
        gym_id: UUID,
        fecha_inicio: Optional[date] = None,
        fecha_fin: Optional[date] = None,
    ) -> ReporteAsistenciaResponse:
        hoy = today_local()
        if fecha_fin is None:
            fecha_fin = hoy
        if fecha_inicio is None:
            fecha_inicio = fecha_fin - timedelta(days=7)

        if fecha_inicio > fecha_fin:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={"codigo": "RANGO_FECHAS_INVALIDO", "mensaje": "La fecha de inicio no puede ser posterior a la fecha de fin"}
            )

        inicio_utc = get_local_day_range_utc(fecha_inicio)[0]
        fin_utc = get_local_day_range_utc(fecha_fin)[1]

        # 1. Totales de accesos
        q_totales = text("""
            SELECT 
                COUNT(*) as total_accesos,
                COUNT(*) FILTER (WHERE resultado IN ('abrio', 'alerta_mora')) as permitidos,
                COUNT(*) FILTER (WHERE resultado = 'negado') as denegados,
                COUNT(*) FILTER (WHERE tipo = 'cortesia') as cortesias
            FROM platform.checkins
            WHERE gimnasio_id = :gym_id AND ts_utc >= :inicio_utc AND ts_utc <= :fin_utc
        """)
        res_tot = await session.execute(q_totales, {"gym_id": gym_id, "inicio_utc": inicio_utc, "fin_utc": fin_utc})
        row_tot = res_tot.mappings().first()

        # 2. Afluencia por hora (00:00 a 23:00 en hora local Bogota)
        q_horas = text("""
            SELECT 
                EXTRACT(HOUR FROM ts_utc AT TIME ZONE 'America/Bogota')::int as hora,
                COUNT(*) as cantidad
            FROM platform.checkins
            WHERE gimnasio_id = :gym_id AND ts_utc >= :inicio_utc AND ts_utc <= :fin_utc
            GROUP BY hora
            ORDER BY hora
        """)
        res_horas = await session.execute(q_horas, {"gym_id": gym_id, "inicio_utc": inicio_utc, "fin_utc": fin_utc})
        horas_dict = {r["hora"]: r["cantidad"] for r in res_horas.mappings().all()}
        afluencia_hora = [
            AfluenciaHoraDto(hora=h, cantidad=horas_dict.get(h, 0))
            for h in range(24)
        ]

        # 3. Afluencia por día de la semana (0=Domingo, 1=Lunes, ..., 6=Sábado)
        nombres_dias = {
            0: "Domingo", 1: "Lunes", 2: "Martes", 3: "Miércoles",
            4: "Jueves", 5: "Viernes", 6: "Sábado"
        }
        q_dias = text("""
            SELECT 
                EXTRACT(DOW FROM ts_utc AT TIME ZONE 'America/Bogota')::int as dia_num,
                COUNT(*) as cantidad
            FROM platform.checkins
            WHERE gimnasio_id = :gym_id AND ts_utc >= :inicio_utc AND ts_utc <= :fin_utc
            GROUP BY dia_num
            ORDER BY dia_num
        """)
        res_dias = await session.execute(q_dias, {"gym_id": gym_id, "inicio_utc": inicio_utc, "fin_utc": fin_utc})
        dias_dict = {r["dia_num"]: r["cantidad"] for r in res_dias.mappings().all()}
        # Ordenar desde Lunes (1) hasta Domingo (0)
        orden_semana = [1, 2, 3, 4, 5, 6, 0]
        afluencia_dia = [
            AfluenciaDiaSemanaDto(dia_num=d, dia_nombre=nombres_dias[d], cantidad=dias_dict.get(d, 0))
            for d in orden_semana
        ]

        # 4. Top 10 deportistas con mayor asistencia
        q_top = text("""
            SELECT 
                c.deportista_id,
                d.nombre,
                d.documento,
                COUNT(*) as total_asistencias,
                MAX(c.ts_utc) as ultima_asistencia
            FROM platform.checkins c
            JOIN platform.deportistas d ON d.id = c.deportista_id
            WHERE c.gimnasio_id = :gym_id 
              AND c.ts_utc >= :inicio_utc AND c.ts_utc <= :fin_utc
              AND c.resultado IN ('abrio', 'alerta_mora')
              AND c.deportista_id IS NOT NULL
            GROUP BY c.deportista_id, d.nombre, d.documento
            ORDER BY total_asistencias DESC
            LIMIT 10
        """)
        res_top = await session.execute(q_top, {"gym_id": gym_id, "inicio_utc": inicio_utc, "fin_utc": fin_utc})
        top_deportistas = [
            TopDeportistaAsistenciaDto(
                deportista_id=r["deportista_id"],
                nombre=r["nombre"],
                documento=r["documento"],
                total_asistencias=r["total_asistencias"],
                ultima_asistencia=r["ultima_asistencia"]
            )
            for r in res_top.mappings().all()
        ]

        return ReporteAsistenciaResponse(
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_fin,
            total_accesos=row_tot["total_accesos"] if row_tot else 0,
            accesos_permitidos=row_tot["permitidos"] if row_tot else 0,
            accesos_denegados=row_tot["denegados"] if row_tot else 0,
            cortesias=row_tot["cortesias"] if row_tot else 0,
            afluencia_por_hora=afluencia_hora,
            afluencia_por_dia_semana=afluencia_dia,
            top_deportistas=top_deportistas
        )

    # =========================================================================
    # 4. RF-43 (Extensión): REPORTE DE DEPORTISTAS INACTIVOS (RETENCIÓN)
    # =========================================================================

    @classmethod
    async def obtener_deportistas_inactivos(
        cls,
        session: AsyncSession,
        gym_id: UUID,
        dias_sin_asistencia: int = 15,
        pagina: int = 1,
        limite: int = 50,
    ) -> ReporteDeportistasInactivosResponse:
        hoy = today_local()
        dias_sin_asistencia = max(1, dias_sin_asistencia)
        pagina = max(1, pagina)
        limite = max(1, min(limite, 10000))

        corte_dt = now_utc() - timedelta(days=dias_sin_asistencia)

        # 1. Configuración de tenant para MembresiaDomainService
        q_conf = text("SELECT dias_gracia_mora, dias_umbral_por_vencer FROM platform.tenant WHERE id = :gym_id")
        res_conf = await session.execute(q_conf, {"gym_id": gym_id})
        row_conf = res_conf.mappings().first()
        gracia = row_conf["dias_gracia_mora"] if row_conf else 3
        umbral = row_conf["dias_umbral_por_vencer"] if row_conf else 5

        # 2. Consultar deportistas activos con membresía vigente
        q_deps = text("""
            SELECT 
                d.id as deportista_id,
                d.nombre as deportista_nombre,
                d.documento as deportista_documento,
                d.correo as deportista_correo,
                d.telefono as deportista_telefono,
                d.activo as deportista_activo,
                m.id as id,
                p.nombre as plan_nombre,
                m.fecha_inicio,
                m.fecha_vencimiento,
                m.cancelada,
                (
                    SELECT MAX(c.ts_utc)
                    FROM platform.checkins c
                    WHERE c.gimnasio_id = :gym_id 
                      AND c.deportista_id = d.id 
                      AND c.resultado IN ('abrio', 'alerta_mora')
                ) as ultimo_checkin_utc
            FROM platform.deportistas d
            JOIN platform.membresias m ON m.deportista_id = d.id
            JOIN platform.planes p ON p.id = m.plan_id
            WHERE d.gimnasio_id = :gym_id 
              AND d.deleted_at IS NULL 
              AND d.activo = true 
              AND m.cancelada = false
              AND m.fecha_vencimiento >= :hoy
            ORDER BY d.nombre ASC
        """)
        res_deps = await session.execute(q_deps, {"gym_id": gym_id, "hoy": hoy})
        raw_items = [dict(r) for r in res_deps.mappings().all()]

        # 3. Evaluar estado canónico de membresía con MembresiaDomainService
        enriquecidos = await MembresiaDomainService.enriquecer_membresias_batch(
            session=session,
            gym_id=gym_id,
            items=raw_items,
            dias_gracia_mora=gracia,
            dias_umbral_por_vencer=umbral
        )

        inactivos: List[DeportistaInactivoItemDto] = []
        now_ts = now_utc()

        for it in enriquecidos:
            estado_calc = it.get("estado_calculado", "sin_membresia")
            # Solo consideramos deportistas con derecho activo a entrenar
            if estado_calc not in ("activo", "por_vencer", "mora"):
                continue

            ult_checkin = it.get("ultimo_checkin_utc")

            if ult_checkin is None:
                # Nunca ha registrado asistencia desde que inició su membresía
                dias_sin_asistir = (hoy - it["fecha_inicio"]).days
                if dias_sin_asistir >= dias_sin_asistencia:
                    inactivos.append(DeportistaInactivoItemDto(
                        deportista_id=it["deportista_id"],
                        nombre=it["deportista_nombre"],
                        documento=it["deportista_documento"],
                        correo=it.get("deportista_correo"),
                        telefono=it.get("deportista_telefono"),
                        plan_nombre=it["plan_nombre"],
                        fecha_vencimiento=it["fecha_vencimiento"],
                        estado_membresia=estado_calc,
                        ultimo_checkin_utc=None,
                        dias_sin_asistir=max(dias_sin_asistir, dias_sin_asistencia)
                    ))
            elif ult_checkin < corte_dt:
                # Asistió alguna vez pero hace más de N días
                dias_sin_asistir = (now_ts - ult_checkin).days
                inactivos.append(DeportistaInactivoItemDto(
                    deportista_id=it["deportista_id"],
                    nombre=it["deportista_nombre"],
                    documento=it["deportista_documento"],
                    correo=it.get("deportista_correo"),
                    telefono=it.get("deportista_telefono"),
                    plan_nombre=it["plan_nombre"],
                    fecha_vencimiento=it["fecha_vencimiento"],
                    estado_membresia=estado_calc,
                    ultimo_checkin_utc=ult_checkin,
                    dias_sin_asistir=dias_sin_asistir
                ))

        # Ordenar por mayor cantidad de días sin asistir (prioridad de contacto)
        inactivos.sort(key=lambda x: x.dias_sin_asistir, reverse=True)

        total = len(inactivos)
        total_paginas = max(1, math.ceil(total / limite))
        offset = (pagina - 1) * limite
        pagina_items = inactivos[offset : offset + limite]

        return ReporteDeportistasInactivosResponse(
            dias_sin_asistencia_umbral=dias_sin_asistencia,
            total=total,
            pagina=pagina,
            limite=limite,
            total_paginas=total_paginas,
            items=pagina_items
        )

    # =========================================================================
    # 5. RF-44: EXPORTACIÓN EN MEMORIA A CSV Y EXCEL (OPENPYXL)
    # =========================================================================

    @classmethod
    async def exportar_reporte(
        cls,
        session: AsyncSession,
        gym_id: UUID,
        tipo_reporte: str,
        formato: str,
        fecha_inicio: Optional[date] = None,
        fecha_fin: Optional[date] = None,
        metodo: Optional[str] = None,
        fuente: Optional[str] = None,
        dias: Optional[int] = None,
        dias_sin_asistencia: int = 15,
    ) -> Response:
        """
        Exporta el reporte solicitado reutilizando EXACTAMENTE el mismo método de agregación.
        CERO divergencia en reglas y CERO archivos temporales en disco (100% en memoria).
        """
        formato = formato.lower().strip()
        tipo_reporte = tipo_reporte.lower().strip()

        if formato not in ("csv", "excel"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={"codigo": "FORMATO_NO_SOPORTADO", "mensaje": "El formato de exportación debe ser 'csv' o 'excel'"}
            )

        headers_cols: List[str] = []
        rows_data: List[List[Any]] = []
        titulo_hoja = tipo_reporte

        if tipo_reporte == "ingresos":
            reporte_ing = await cls.obtener_reporte_ingresos(
                session=session,
                gym_id=gym_id,
                fecha_inicio=fecha_inicio,
                fecha_fin=fecha_fin,
                metodo=metodo,
                fuente=fuente
            )
            headers_cols = ["ID", "Fecha/Hora", "Tipo", "Descripción", "Monto", "Método", "Tipo Medio", "Anulado", "Motivo Anulación", "Registrado Por"]
            for t in reporte_ing.transacciones:
                rows_data.append([
                    str(t.id),
                    t.fecha_hora.strftime("%Y-%m-%d %H:%M:%S"),
                    t.tipo,
                    t.descripcion,
                    float(t.monto),
                    t.metodo,
                    t.tipo_medio,
                    "SÍ" if t.anulado else "NO",
                    t.motivo_anulacion or "",
                    t.registrado_por_nombre or ""
                ])
            titulo_hoja = "Ingresos"

        elif tipo_reporte == "membresias_por_vencer":
            reporte_venc = await cls.obtener_membresias_por_vencer(
                session=session,
                gym_id=gym_id,
                dias=dias,
                pagina=1,
                limite=10000
            )
            headers_cols = ["Nombre Deportista", "Documento", "Correo", "Teléfono", "Plan", "Fecha Inicio", "Fecha Vencimiento", "Días Restantes", "Congelada"]
            for it in reporte_venc.items:
                rows_data.append([
                    it.nombre,
                    it.documento,
                    it.correo or "",
                    it.telefono or "",
                    it.plan_nombre,
                    it.fecha_inicio.strftime("%Y-%m-%d"),
                    it.fecha_vencimiento.strftime("%Y-%m-%d"),
                    it.dias_restantes,
                    "SÍ" if it.congelamiento_activo else "NO"
                ])
            titulo_hoja = "Membresías Por Vencer"

        elif tipo_reporte == "asistencia":
            reporte_asist = await cls.obtener_reporte_asistencia(
                session=session,
                gym_id=gym_id,
                fecha_inicio=fecha_inicio,
                fecha_fin=fecha_fin
            )
            headers_cols = ["Hora Local", "Cantidad Accesos"]
            for ah in reporte_asist.afluencia_por_hora:
                rows_data.append([f"{ah.hora:02d}:00", ah.cantidad])
            titulo_hoja = "Afluencia Horaria"

        elif tipo_reporte == "deportistas_inactivos":
            reporte_inac = await cls.obtener_deportistas_inactivos(
                session=session,
                gym_id=gym_id,
                dias_sin_asistencia=dias_sin_asistencia,
                pagina=1,
                limite=10000
            )
            headers_cols = ["Nombre", "Documento", "Correo", "Teléfono", "Plan", "Fecha Vencimiento", "Estado Membresía", "Último Check-in", "Días Sin Asistir"]
            for it in reporte_inac.items:
                rows_data.append([
                    it.nombre,
                    it.documento,
                    it.correo or "",
                    it.telefono or "",
                    it.plan_nombre,
                    it.fecha_vencimiento.strftime("%Y-%m-%d"),
                    it.estado_membresia,
                    it.ultimo_checkin_utc.strftime("%Y-%m-%d %H:%M:%S") if it.ultimo_checkin_utc else "Sin registro",
                    it.dias_sin_asistir
                ])
            titulo_hoja = "Deportistas Inactivos"

        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={"codigo": "TIPO_REPORTE_INVALIDO", "mensaje": f"Tipo de reporte desconocido: '{tipo_reporte}'"}
            )

        timestamp_str = now_local().strftime("%Y%m%d_%H%M%S")

        # Serialización en CSV con UTF-8 BOM
        if formato == "csv":
            str_io = io.StringIO()
            # BOM para que Excel en Windows decodifique UTF-8 automáticamente
            str_io.write("\ufeff")
            writer = csv.writer(str_io, delimiter=",", quoting=csv.QUOTE_MINIMAL)
            writer.writerow(headers_cols)
            for r in rows_data:
                writer.writerow(r)

            content_bytes = str_io.getvalue().encode("utf-8")
            filename = f"reporte_{tipo_reporte}_{timestamp_str}.csv"
            return Response(
                content=content_bytes,
                media_type="text/csv; charset=utf-8",
                headers={"Content-Disposition": f'attachment; filename="{filename}"'}
            )

        # Serialización en Excel (.xlsx) con openpyxl
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = titulo_hoja[:30]

            # Estilo cabeceras
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

            ws.append(headers_cols)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for r in rows_data:
                ws.append(r)

            # Auto-ancho de columnas
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            byte_io = io.BytesIO()
            wb.save(byte_io)
            byte_io.seek(0)
            content_bytes = byte_io.getvalue()

            filename = f"reporte_{tipo_reporte}_{timestamp_str}.xlsx"
            return Response(
                content=content_bytes,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f'attachment; filename="{filename}"'}
            )
