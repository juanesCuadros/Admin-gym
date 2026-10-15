"""
GW-RF-55 · Carga inicial de clientes (VISION §4.6, contrato 2.4).
"""
import io
import time
from datetime import date, datetime, timedelta
from uuid import uuid4

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from app.core.database import async_session_maker
from app.core.timezone import today_local
from tests.conftest import Fabrica, login

COLUMNAS = "tipo_documento;documento;nombre;telefono;correo;fecha_nacimiento;plan_nombre;fecha_vencimiento"


def _f(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def _csv(*filas: str, encabezado: str = COLUMNAS) -> bytes:
    return ("﻿" + "\r\n".join((encabezado,) + filas) + "\r\n").encode("utf-8")


async def _preparar(cliente: AsyncClient, fabrica: Fabrica, rol: str = "jefe"):
    gym = await fabrica.crear_gimnasio()
    staff = await fabrica.crear_staff(gym, rol=rol)
    tok = (await login(cliente, gym, staff)).json()["access_token"]
    async with async_session_maker() as s:
        await s.execute(text("SELECT set_config('app.gimnasio_id', :g, true)"), {"g": str(gym.id)})
        await s.execute(
            text("""
                INSERT INTO platform.planes (gimnasio_id, nombre, precio, duracion_dias, tipo)
                VALUES (:g, 'Mensual', 80000, 30, 'individual'), (:g, 'Trimestral', 210000, 90, 'individual')
            """),
            {"g": gym.id},
        )
        await s.commit()
    return gym, {"Authorization": f"Bearer {tok}"}


async def _contar(fabrica: Fabrica, gym, sql: str) -> int:
    async with fabrica.admin_engine.connect() as conn:
        return (await conn.execute(text(sql), {"g": gym.id})).scalar_one()


def _archivo(contenido: bytes, nombre: str = "clientes.csv"):
    tipo = "text/csv" if nombre.endswith(".csv") else "application/octet-stream"
    return {"archivo": (nombre, contenido, tipo)}


async def _aplicar(cliente, headers, contenido, nombre="clientes.csv", declara=True):
    return await cliente.post(
        "/api/v1/carga-inicial/aplicar",
        files=_archivo(contenido, nombre),
        data={"declara_consentimientos": "true" if declara else "false"},
        headers=headers,
    )


@pytest.mark.asyncio
async def test_plantilla_csv_y_xlsx(cliente: AsyncClient, fabrica: Fabrica):
    gym, headers = await _preparar(cliente, fabrica)

    r_csv = await cliente.get("/api/v1/carga-inicial/plantilla?formato=csv", headers=headers)
    assert r_csv.status_code == 200
    assert r_csv.content.startswith("﻿".encode("utf-8"))
    assert r_csv.content.decode("utf-8-sig").strip() == COLUMNAS

    r_xlsx = await cliente.get("/api/v1/carga-inicial/plantilla?formato=xlsx", headers=headers)
    assert r_xlsx.status_code == 200
    hoja = load_workbook(io.BytesIO(r_xlsx.content)).worksheets[0]
    assert [c.value for c in hoja[1]] == COLUMNAS.split(";")


@pytest.mark.asyncio
async def test_simular_reporta_errores_por_fila_sin_guardar(cliente: AsyncClient, fabrica: Fabrica):
    gym, headers = await _preparar(cliente, fabrica)
    hoy = today_local()
    menor = date(hoy.year - 15, 1, 1)
    contenido = _csv(
        f"CC;1001;Ana Válida;3001112233;ana@correo.com;01/02/1990;Mensual;{_f(hoy + timedelta(days=20))}",
        f"CC;1002;Plan Raro;;;;Anual Plus;{_f(hoy + timedelta(days=20))}",
        "CC;1003;Fecha Mala;;;;Mensual;31/02/2026",
        f"TI;1004;Menor de Edad;;;{_f(menor)};;",
        "XX;1005;Tipo Malo;;;;;",
        "CC;1006;Sin Vencimiento;;;;Mensual;",
        "CC;1007;Correo Malo;;no-es-correo;;;",
    )
    r = await cliente.post("/api/v1/carga-inicial/simular", files=_archivo(contenido), headers=headers)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["filas_validas"] == 1
    assert body["filas_con_error"] == 6
    codigos = {(e["fila"], e["codigo"]) for e in body["errores"]}
    assert (3, "PLAN_NO_EXISTE") in codigos
    assert (4, "FECHA_INVALIDA") in codigos
    assert (5, "MENOR_DE_EDAD") in codigos
    assert (6, "TIPO_DOCUMENTO_INVALIDO") in codigos
    assert (7, "VENCIMIENTO_REQUERIDO") in codigos
    assert (8, "CORREO_INVALIDO") in codigos

    assert await _contar(fabrica, gym, "SELECT count(*) FROM platform.deportistas WHERE gimnasio_id = :g") == 0


@pytest.mark.asyncio
async def test_documento_repetido_no_carga_nada_ca1(cliente: AsyncClient, fabrica: Fabrica):
    """CA1: un archivo con un documento repetido no carga nada y reporta la fila."""
    gym, headers = await _preparar(cliente, fabrica)
    venc = _f(today_local() + timedelta(days=15))
    contenido = _csv(
        f"CC;2001;Primero;;;;Mensual;{venc}",
        f"CC;2002;Segundo;;;;Mensual;{venc}",
        f"CC;2.001;Repetido con puntos;;;;Mensual;{venc}",
    )
    r = await _aplicar(cliente, headers, contenido)
    assert r.status_code == 422, r.text
    error = r.json()["error"]
    assert error["codigo"] == "CARGA_CON_ERRORES"
    assert error["detalles"]["filas_con_error"] == 1
    [e] = error["detalles"]["errores"]
    assert (e["fila"], e["columna"], e["codigo"]) == (4, "documento", "DOCUMENTO_REPETIDO")
    assert "fila 2" in e["mensaje"]

    assert await _contar(fabrica, gym, "SELECT count(*) FROM platform.deportistas WHERE gimnasio_id = :g") == 0
    assert await _contar(fabrica, gym, "SELECT count(*) FROM platform.membresias WHERE gimnasio_id = :g") == 0


@pytest.mark.asyncio
async def test_documento_ya_registrado_en_el_gimnasio(cliente: AsyncClient, fabrica: Fabrica):
    gym, headers = await _preparar(cliente, fabrica)
    r0 = await cliente.post(
        "/api/v1/deportistas",
        json={"tipo_documento": "CC", "documento": "3001", "nombre": "Ya Existe", "consentimiento_1581": True},
        headers=headers,
    )
    assert r0.status_code == 201, r0.text

    r = await _aplicar(cliente, headers, _csv("CC;3001;Ya Existe Otra Vez;;;;;", "CE;3001;Otro Tipo;;;;;"))
    assert r.status_code == 422
    [e] = r.json()["error"]["detalles"]["errores"]
    assert (e["fila"], e["codigo"]) == (2, "DOCUMENTO_YA_REGISTRADO")


@pytest.mark.asyncio
async def test_estado_de_migrados_activo_y_en_gracia_ca2(cliente: AsyncClient, fabrica: Fabrica):
    """CA2: vence en 10 días → activo; venció ayer con gracia 3 → en_gracia."""
    gym, headers = await _preparar(cliente, fabrica)
    hoy = today_local()
    contenido = _csv(
        f"CC;4001;Cliente Activo;;;;Mensual;{_f(hoy + timedelta(days=10))}",
        f"CC;4002;Cliente En Gracia;;;;mensual;{_f(hoy - timedelta(days=1))}",
        "CC;4003;Cliente Sin Plan;;;;;",
    )
    r = await _aplicar(cliente, headers, contenido)
    assert r.status_code == 201, r.text
    assert r.json() == {"deportistas_creados": 3, "membresias_creadas": 2}

    esperados = {"4001": "activo", "4002": "en_gracia", "4003": "sin_membresia"}
    for doc, estado in esperados.items():
        rb = await cliente.get(f"/api/v1/deportistas/buscar?q={doc}", headers=headers)
        assert rb.status_code == 200, rb.text
        [dep] = rb.json()
        assert dep["estado"] == estado, (doc, dep)

    async with fabrica.admin_engine.connect() as conn:
        filas = (await conn.execute(
            text("""
                SELECT m.origen, m.fecha_inicio, m.fecha_vencimiento, m.plan_duracion_dias_vendida,
                       d.consentimiento_1581, d.consentimiento_version
                FROM platform.membresias m JOIN platform.deportistas d ON d.id = m.deportista_id
                WHERE m.gimnasio_id = :g
            """),
            {"g": gym.id},
        )).mappings().all()
    assert len(filas) == 2
    for f in filas:
        assert f["origen"] == "migracion"
        assert f["fecha_inicio"] + timedelta(days=f["plan_duracion_dias_vendida"]) == f["fecha_vencimiento"]
        assert f["consentimiento_1581"] is True
        assert f["consentimiento_version"] == "migracion"


@pytest.mark.asyncio
async def test_migradas_no_cuentan_en_ingresos_ca3(cliente: AsyncClient, fabrica: Fabrica):
    """CA3: los ingresos del día no cuentan las membresías migradas."""
    gym, headers = await _preparar(cliente, fabrica)
    hoy = today_local()
    url = f"/api/v1/reportes/ingresos?fecha_inicio={hoy}&fecha_fin={hoy}"

    antes = await cliente.get(url, headers=headers)
    assert antes.status_code == 200, antes.text

    venc = _f(hoy + timedelta(days=25))
    r = await _aplicar(cliente, headers, _csv(f"CC;5001;Uno;;;;Mensual;{venc}", f"CC;5002;Dos;;;;Trimestral;{venc}"))
    assert r.status_code == 201, r.text

    despues = await cliente.get(url, headers=headers)
    assert despues.status_code == 200
    assert float(despues.json()["total_ingresos_bruto"]) == float(antes.json()["total_ingresos_bruto"]) == 0
    assert despues.json()["transacciones"] == []
    assert await _contar(fabrica, gym, "SELECT count(*) FROM platform.pagos_membresia WHERE gimnasio_id = :g") == 0


@pytest.mark.asyncio
async def test_exige_declaracion_de_consentimiento(cliente: AsyncClient, fabrica: Fabrica):
    gym, headers = await _preparar(cliente, fabrica)
    r = await _aplicar(cliente, headers, _csv("CC;6001;Sin Declarar;;;;;"), declara=False)
    assert r.status_code == 422
    assert r.json()["error"]["codigo"] == "DECLARACION_CONSENTIMIENTO_REQUERIDA"
    assert await _contar(fabrica, gym, "SELECT count(*) FROM platform.deportistas WHERE gimnasio_id = :g") == 0


@pytest.mark.asyncio
async def test_solo_el_jefe(cliente: AsyncClient, fabrica: Fabrica):
    gym, headers = await _preparar(cliente, fabrica, rol="recepcionista")
    r_sim = await cliente.post("/api/v1/carga-inicial/simular", files=_archivo(_csv("CC;7001;X Y;;;;;")), headers=headers)
    r_apl = await _aplicar(cliente, headers, _csv("CC;7001;X Y;;;;;"))
    r_pla = await cliente.get("/api/v1/carga-inicial/plantilla", headers=headers)
    assert r_sim.status_code == r_apl.status_code == r_pla.status_code == 403


@pytest.mark.asyncio
async def test_excel_y_csv_con_coma(cliente: AsyncClient, fabrica: Fabrica):
    gym, headers = await _preparar(cliente, fabrica)
    hoy = today_local()

    libro = Workbook()
    hoja = libro.active
    hoja.append(COLUMNAS.split(";"))
    # Excel guarda el documento como número y las fechas como fecha.
    hoja.append(["cc", 8001, "Desde Excel", 3001234567, None, datetime(1990, 5, 4), "MENSUAL", hoy + timedelta(days=12)])
    salida = io.BytesIO()
    libro.save(salida)
    r_x = await _aplicar(cliente, headers, salida.getvalue(), nombre="clientes.xlsx")
    assert r_x.status_code == 201, r_x.text

    contenido = _csv(
        f"CC,1.020.304,Desde CSV,,,,Mensual,{(hoy + timedelta(days=12)).isoformat()}",
        encabezado=COLUMNAS.replace(";", ","),
    )
    r_c = await _aplicar(cliente, headers, contenido)
    assert r_c.status_code == 201, r_c.text

    async with fabrica.admin_engine.connect() as conn:
        docs = (await conn.execute(
            text("SELECT documento, tipo_documento, fecha_nacimiento FROM platform.deportistas WHERE gimnasio_id = :g ORDER BY documento"),
            {"g": gym.id},
        )).all()
    assert [(d[0], d[1]) for d in docs] == [("1020304", "CC"), ("8001", "CC")]
    assert docs[1][2] == date(1990, 5, 4)


@pytest.mark.asyncio
async def test_archivo_invalido(cliente: AsyncClient, fabrica: Fabrica):
    gym, headers = await _preparar(cliente, fabrica)
    r_fmt = await cliente.post("/api/v1/carga-inicial/simular", files=_archivo(b"x", "clientes.pdf"), headers=headers)
    assert r_fmt.json()["error"]["codigo"] == "FORMATO_NO_SOPORTADO"

    r_cols = await cliente.post(
        "/api/v1/carga-inicial/simular",
        files=_archivo(_csv("CC;1;Ana", encabezado="tipo_documento;documento;nombre")),
        headers=headers,
    )
    assert r_cols.status_code == 422
    assert r_cols.json()["error"]["codigo"] == "PLANTILLA_INVALIDA"
    assert "plan_nombre" in r_cols.json()["error"]["detalles"]["faltantes"]

    r_vacio = await cliente.post("/api/v1/carga-inicial/simular", files=_archivo(_csv()), headers=headers)
    assert r_vacio.json()["error"]["codigo"] == "ARCHIVO_VACIO"


@pytest.mark.asyncio
async def test_respeta_limite_del_plan_saas(cliente: AsyncClient, fabrica: Fabrica):
    gym, headers = await _preparar(cliente, fabrica)
    plan_id = uuid4()
    async with fabrica.admin_engine.begin() as conn:
        await conn.execute(
            text("""
                INSERT INTO superadmin.planes_suscripcion
                  (id, codigo, nombre, tipo, precio, duracion_meses, max_clientes_activos, max_staff, activo, publico)
                VALUES (:id, :cod, 'Plan Micro Carga', 'pago', 50000, 1, 2, 5, true, true)
            """),
            {"id": plan_id, "cod": f"carga_{uuid4().hex[:6]}"},
        )
        await conn.execute(
            text("""
                INSERT INTO superadmin.suscripciones (gimnasio_id, valor_mensual, tipo_inicio, vigente_desde, plan_id)
                VALUES (:g, 50000, 'cliente_activo', current_date - 1, :p)
            """),
            {"g": gym.superadmin_id, "p": plan_id},
        )

    r = await _aplicar(cliente, headers, _csv("CC;9001;Uno Uno;;;;;", "CC;9002;Dos Dos;;;;;", "CC;9003;Tres Tres;;;;;"))
    assert r.status_code == 409, r.text
    detalles = r.json()["error"]["detalles"]
    assert r.json()["error"]["codigo"] == "LIMITE_PLAN_CLIENTES"
    assert (detalles["limite"], detalles["actuales"], detalles["a_cargar"]) == (2, 0, 3)
    assert await _contar(fabrica, gym, "SELECT count(*) FROM platform.deportistas WHERE gimnasio_id = :g") == 0


@pytest.mark.asyncio
async def test_auditoria_sin_datos_personales(cliente: AsyncClient, fabrica: Fabrica):
    gym, headers = await _preparar(cliente, fabrica)
    r = await _aplicar(cliente, headers, _csv("CC;9101;Persona Auditada;3009998877;auditada@correo.com;;;"))
    assert r.status_code == 201, r.text
    async with fabrica.admin_engine.connect() as conn:
        [fila] = (await conn.execute(
            text("SELECT detalle::text FROM platform.auditoria_gym WHERE gimnasio_id = :g AND accion = 'CARGA_INICIAL'"),
            {"g": gym.id},
        )).all()
    for dato in ("9101", "Persona", "3009998877", "auditada"):
        assert dato not in fila[0]
    assert '"deportistas_creados": 1' in fila[0]


@pytest.mark.asyncio
async def test_dos_mil_filas_en_menos_de_30_segundos(cliente: AsyncClient, fabrica: Fabrica):
    gym, headers = await _preparar(cliente, fabrica)
    venc = _f(today_local() + timedelta(days=20))
    filas = [f"CC;{10_000_000 + i};Cliente Masivo {i};300{i:07d};c{i}@correo.com;01/01/1990;Mensual;{venc}" for i in range(2000)]
    inicio = time.monotonic()
    r = await _aplicar(cliente, headers, _csv(*filas))
    duracion = time.monotonic() - inicio
    assert r.status_code == 201, r.text
    assert r.json() == {"deportistas_creados": 2000, "membresias_creadas": 2000}
    assert duracion < 30, f"La carga tardó {duracion:.1f} s"
