"""
Carga inicial de clientes y membresías vigentes (GW-RF-55, VISION §4.6, contrato 2.4).

- El gimnasio llena una plantilla CSV o Excel con sus clientes actuales.
- `simular` valida todo y devuelve los errores por fila sin guardar nada.
- `aplicar` repite la validación y guarda todo o nada en la transacción del request.
- Las membresías quedan con `origen = 'migracion'` y sin pago ni movimiento de caja, así que no
  alteran el cuadre ni los reportes de ingresos.
- El Jefe declara que el gimnasio conserva la autorización de datos (Ley 1581) de estos clientes;
  el consentimiento queda con versión 'migracion'.
- Los menores de edad no se cargan: necesitan acudiente y se registran uno a uno.
"""
import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from typing import Any, Dict, List, Optional, Tuple
from uuid import UUID, uuid4

from email_validator import EmailNotValidError, validate_email
from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.audit import AuditService
from app.core.timezone import today_local
from app.modules.carga_inicial.schemas import (
    AplicacionCargaResponse,
    ErrorFilaDto,
    SimulacionCargaResponse,
)

COLUMNAS = [
    "tipo_documento",
    "documento",
    "nombre",
    "telefono",
    "correo",
    "fecha_nacimiento",
    "plan_nombre",
    "fecha_vencimiento",
]
TIPOS_DOCUMENTO = ("CC", "TI", "CE", "PAS", "PPT", "RC")
MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 5000
CONSENTIMIENTO_VERSION_MIGRACION = "migracion"
FORMATOS_FECHA = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y")


@dataclass
class _FilaValida:
    tipo_documento: str
    documento: str
    nombre: str
    telefono: Optional[str]
    correo: Optional[str]
    fecha_nacimiento: Optional[date]
    plan: Optional[Dict[str, Any]]
    fecha_vencimiento: Optional[date]


def _error(status_code: int, codigo: str, mensaje: str, detalles: Optional[dict] = None) -> HTTPException:
    detail: Dict[str, Any] = {"codigo": codigo, "mensaje": mensaje}
    if detalles is not None:
        detail["detalles"] = detalles
    return HTTPException(status_code=status_code, detail=detail)


# =====================================================================
# Lectura del archivo
# =====================================================================

def _texto(valor: Any) -> str:
    """Normaliza una celda a texto. Excel entrega los documentos numéricos como int o float."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    return str(valor).strip()


def _filas_csv(contenido: bytes) -> List[List[Any]]:
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel en Windows guarda "CSV" en la página de códigos del sistema.
        texto = contenido.decode("cp1252", errors="replace")
    primera = next((l for l in texto.splitlines() if l.strip()), "")
    # Excel con configuración regional de Colombia separa con punto y coma.
    separador = ";" if primera.count(";") >= primera.count(",") else ","
    return [fila for fila in csv.reader(io.StringIO(texto), delimiter=separador)]


def _filas_xlsx(contenido: bytes) -> List[List[Any]]:
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise _error(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "ARCHIVO_INVALIDO",
            "No se pudo leer el archivo de Excel. Verifica que sea un .xlsx sin contraseña.",
        )
    try:
        hoja = libro.worksheets[0]
        return [list(fila) for fila in hoja.iter_rows(values_only=True)]
    finally:
        libro.close()


def leer_archivo(nombre_archivo: Optional[str], contenido: bytes) -> List[Tuple[int, Dict[str, Any]]]:
    """Devuelve [(numero_de_fila, {columna: valor})] sin las filas vacías."""
    if len(contenido) > MAX_BYTES:
        raise _error(
            status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            "ARCHIVO_DEMASIADO_GRANDE",
            "El archivo supera el tamaño máximo de 5 MB.",
        )
    nombre = (nombre_archivo or "").lower()
    if nombre.endswith(".csv"):
        filas = _filas_csv(contenido)
    elif nombre.endswith(".xlsx"):
        filas = _filas_xlsx(contenido)
    else:
        raise _error(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "FORMATO_NO_SOPORTADO",
            "El archivo debe ser la plantilla en formato .csv o .xlsx.",
        )

    # El encabezado es la primera fila con contenido.
    inicio = next((i for i, f in enumerate(filas) if any(_texto(c) for c in f)), None)
    if inicio is None:
        raise _error(status.HTTP_422_UNPROCESSABLE_ENTITY, "ARCHIVO_VACIO", "El archivo no tiene datos.")

    encabezado = [_texto(c).lstrip("﻿").strip().lower() for c in filas[inicio]]
    faltantes = [c for c in COLUMNAS if c not in encabezado]
    if faltantes:
        raise _error(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "PLANTILLA_INVALIDA",
            "Al archivo le faltan columnas de la plantilla: " + ", ".join(faltantes) + ".",
            {"faltantes": faltantes},
        )
    indice = {c: encabezado.index(c) for c in COLUMNAS}

    resultado: List[Tuple[int, Dict[str, Any]]] = []
    for pos in range(inicio + 1, len(filas)):
        fila = filas[pos]
        if not any(_texto(c) for c in fila):
            continue
        valores = {c: (fila[i] if i < len(fila) else None) for c, i in indice.items()}
        resultado.append((pos + 1, valores))
        if len(resultado) > MAX_FILAS:
            raise _error(
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                "DEMASIADAS_FILAS",
                f"El archivo supera el máximo de {MAX_FILAS} clientes por carga. Divídelo en varios archivos.",
            )

    if not resultado:
        raise _error(status.HTTP_422_UNPROCESSABLE_ENTITY, "ARCHIVO_VACIO", "El archivo no tiene clientes para cargar.")
    return resultado


# =====================================================================
# Validación
# =====================================================================

def _clave_plan(nombre: str) -> str:
    sin_tildes = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", sin_tildes).strip().lower()


def _fecha(valor: Any) -> Optional[date]:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = _texto(valor)
    for formato in FORMATOS_FECHA:
        try:
            return datetime.strptime(s, formato).date()
        except ValueError:
            continue
    return None


def _edad(nacimiento: date, hoy: date) -> int:
    return hoy.year - nacimiento.year - ((hoy.month, hoy.day) < (nacimiento.month, nacimiento.day))


async def _validar(
    session: AsyncSession, gym_id: UUID, filas: List[Tuple[int, Dict[str, Any]]]
) -> Tuple[List[_FilaValida], List[ErrorFilaDto], int]:
    """Devuelve (filas válidas, errores, número de filas con error)."""
    hoy = today_local()

    res_planes = await session.execute(
        text("""
            SELECT id, nombre, duracion_dias, precio
            FROM platform.planes
            WHERE gimnasio_id = :gym_id AND activo = true
        """),
        {"gym_id": gym_id},
    )
    planes = {_clave_plan(p["nombre"]): dict(p) for p in res_planes.mappings().all()}

    def _documento(valor: Any) -> str:
        # Se quitan puntos, comas y espacios: "1.020.304" es el mismo documento que "1020304".
        return re.sub(r"[.,\s]", "", _texto(valor)).upper()

    documentos = list({_documento(v["documento"]) for _, v in filas if _documento(v["documento"])})
    existentes: set = set()
    if documentos:
        res_exist = await session.execute(
            text("""
                SELECT tipo_documento, documento FROM platform.deportistas
                WHERE gimnasio_id = :gym_id AND deleted_at IS NULL AND documento = ANY(:docs)
            """),
            {"gym_id": gym_id, "docs": documentos},
        )
        existentes = {(r[0], r[1]) for r in res_exist.fetchall()}

    validas: List[_FilaValida] = []
    errores: List[ErrorFilaDto] = []
    filas_con_error = 0
    vistos: Dict[Tuple[str, str], int] = {}

    for numero, v in filas:
        errs: List[ErrorFilaDto] = []

        def err(columna: str, codigo: str, mensaje: str) -> None:
            errs.append(ErrorFilaDto(fila=numero, columna=columna, codigo=codigo, mensaje=mensaje))

        tipo = _texto(v["tipo_documento"]).upper()
        if tipo not in TIPOS_DOCUMENTO:
            err("tipo_documento", "TIPO_DOCUMENTO_INVALIDO",
                "El tipo de documento debe ser uno de: " + ", ".join(TIPOS_DOCUMENTO) + ".")

        documento = _documento(v["documento"])
        if not documento:
            err("documento", "DOCUMENTO_REQUERIDO", "El documento es obligatorio.")
        elif not (3 <= len(documento) <= 30):
            err("documento", "DOCUMENTO_INVALIDO", "El documento debe tener entre 3 y 30 caracteres.")
        elif tipo in TIPOS_DOCUMENTO:
            clave = (tipo, documento)
            if clave in vistos:
                err("documento", "DOCUMENTO_REPETIDO",
                    f"El documento {tipo} {documento} ya aparece en la fila {vistos[clave]} del archivo.")
            elif clave in existentes:
                err("documento", "DOCUMENTO_YA_REGISTRADO",
                    f"Ya existe un cliente con el documento {tipo} {documento} en el gimnasio.")
            else:
                vistos[clave] = numero

        nombre = re.sub(r"\s+", " ", _texto(v["nombre"]))
        if len(nombre) < 2:
            err("nombre", "NOMBRE_REQUERIDO", "El nombre es obligatorio.")
        elif len(nombre) > 150:
            err("nombre", "NOMBRE_INVALIDO", "El nombre no puede superar 150 caracteres.")

        telefono = _texto(v["telefono"]) or None
        if telefono and len(telefono) > 25:
            err("telefono", "TELEFONO_INVALIDO", "El teléfono no puede superar 25 caracteres.")

        correo = _texto(v["correo"]) or None
        if correo:
            try:
                correo = validate_email(correo, check_deliverability=False).normalized
            except EmailNotValidError:
                err("correo", "CORREO_INVALIDO", "El correo no tiene un formato válido.")

        nacimiento = None
        if _texto(v["fecha_nacimiento"]):
            nacimiento = _fecha(v["fecha_nacimiento"])
            if nacimiento is None:
                err("fecha_nacimiento", "FECHA_INVALIDA", "La fecha de nacimiento debe tener el formato dd/mm/aaaa.")
            elif nacimiento > hoy:
                err("fecha_nacimiento", "FECHA_INVALIDA", "La fecha de nacimiento no puede ser futura.")
            elif _edad(nacimiento, hoy) < 18:
                err("fecha_nacimiento", "MENOR_DE_EDAD",
                    "Es menor de edad: regístralo a mano con los datos y la autorización de su acudiente.")

        plan = None
        vencimiento = None
        tiene_plan = bool(_texto(v["plan_nombre"]))
        tiene_venc = bool(_texto(v["fecha_vencimiento"]))
        if tiene_plan:
            plan = planes.get(_clave_plan(_texto(v["plan_nombre"])))
            if plan is None:
                err("plan_nombre", "PLAN_NO_EXISTE",
                    f"No existe un plan activo llamado «{_texto(v['plan_nombre'])}» en el gimnasio.")
        if tiene_venc:
            vencimiento = _fecha(v["fecha_vencimiento"])
            if vencimiento is None:
                err("fecha_vencimiento", "FECHA_INVALIDA", "La fecha de vencimiento debe tener el formato dd/mm/aaaa.")
        if tiene_plan and not tiene_venc:
            err("fecha_vencimiento", "VENCIMIENTO_REQUERIDO", "Si el cliente tiene plan, indica la fecha de vencimiento.")
        if tiene_venc and not tiene_plan:
            err("plan_nombre", "PLAN_REQUERIDO", "Si el cliente tiene fecha de vencimiento, indica el plan.")

        if errs:
            errores.extend(errs)
            filas_con_error += 1
        else:
            validas.append(_FilaValida(
                tipo_documento=tipo,
                documento=documento,
                nombre=nombre,
                telefono=telefono,
                correo=correo,
                fecha_nacimiento=nacimiento,
                plan=plan,
                fecha_vencimiento=vencimiento,
            ))

    return validas, errores, filas_con_error


async def _validar_limite_plan(session: AsyncSession, gym_id: UUID, nuevos: int) -> None:
    """Igual que el registro uno a uno: la carga no puede superar los clientes activos del plan SaaS."""
    res_plan = await session.execute(
        text("SELECT max_clientes_activos, nombre FROM platform.plan_actual(:gym_id)"),
        {"gym_id": gym_id},
    )
    plan = res_plan.mappings().first()
    if not plan or plan["max_clientes_activos"] is None:
        return
    res_count = await session.execute(
        text("""
            SELECT count(*) FROM platform.deportistas
            WHERE gimnasio_id = :gym_id AND deleted_at IS NULL AND activo = true
        """),
        {"gym_id": gym_id},
    )
    actuales = res_count.scalar_one()
    limite = plan["max_clientes_activos"]
    if actuales + nuevos > limite:
        raise _error(
            status.HTTP_409_CONFLICT,
            "LIMITE_PLAN_CLIENTES",
            f"La carga supera el límite de {limite} clientes activos del plan '{plan['nombre']}': "
            f"ya tienes {actuales} y el archivo trae {nuevos}. Mejora tu plan o reduce el archivo.",
            {"limite": limite, "actuales": actuales, "a_cargar": nuevos, "plan_nombre": plan["nombre"]},
        )


# =====================================================================
# Servicio
# =====================================================================

class CargaInicialService:

    @staticmethod
    def plantilla(formato: str) -> Tuple[bytes, str, str]:
        """Devuelve (contenido, media_type, nombre_de_archivo) de la plantilla vacía."""
        if formato == "csv":
            contenido = ("﻿" + ";".join(COLUMNAS) + "\r\n").encode("utf-8")
            return contenido, "text/csv; charset=utf-8", "plantilla_carga_inicial.csv"

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Clientes"
        hoja.append(COLUMNAS)
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        for letra, ancho in zip("ABCDEFGH", (16, 16, 32, 16, 30, 18, 24, 18)):
            hoja.column_dimensions[letra].width = ancho

        ayuda = libro.create_sheet("Instrucciones")
        ayuda.append(["Columna", "Obligatoria", "Formato"])
        for celda in ayuda[1]:
            celda.font = Font(bold=True)
        for fila in (
            ("tipo_documento", "Sí", ", ".join(TIPOS_DOCUMENTO)),
            ("documento", "Sí", "Sin puntos ni espacios. No puede repetirse."),
            ("nombre", "Sí", "Nombre completo"),
            ("telefono", "No", "Celular o fijo"),
            ("correo", "No", "correo@dominio.com"),
            ("fecha_nacimiento", "No", "dd/mm/aaaa. Los menores de 18 años se registran a mano con su acudiente."),
            ("plan_nombre", "Con vencimiento", "Nombre exacto de un plan activo del gimnasio"),
            ("fecha_vencimiento", "Con plan", "dd/mm/aaaa. Día en que vence la membresía actual."),
        ):
            ayuda.append(list(fila))
        ayuda.column_dimensions["A"].width = 20
        ayuda.column_dimensions["B"].width = 16
        ayuda.column_dimensions["C"].width = 80

        salida = io.BytesIO()
        libro.save(salida)
        return (
            salida.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "plantilla_carga_inicial.xlsx",
        )

    @classmethod
    async def simular(
        cls, session: AsyncSession, gym_id: UUID, nombre_archivo: Optional[str], contenido: bytes
    ) -> SimulacionCargaResponse:
        filas = leer_archivo(nombre_archivo, contenido)
        validas, errores, con_error = await _validar(session, gym_id, filas)
        await _validar_limite_plan(session, gym_id, len(filas))
        return SimulacionCargaResponse(filas_validas=len(validas), filas_con_error=con_error, errores=errores)

    @classmethod
    async def aplicar(
        cls,
        session: AsyncSession,
        gym_id: UUID,
        staff_id: UUID,
        staff_nombre: str,
        nombre_archivo: Optional[str],
        contenido: bytes,
        declara_consentimientos: bool,
    ) -> AplicacionCargaResponse:
        if not declara_consentimientos:
            raise _error(
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                "DECLARACION_CONSENTIMIENTO_REQUERIDA",
                "Para cargar los clientes debes declarar que el gimnasio conserva su autorización de "
                "tratamiento de datos (Ley 1581).",
            )

        filas = leer_archivo(nombre_archivo, contenido)
        validas, errores, con_error = await _validar(session, gym_id, filas)
        if errores:
            raise _error(
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                "CARGA_CON_ERRORES",
                f"No se cargó ningún cliente: {con_error} fila(s) tienen errores.",
                SimulacionCargaResponse(
                    filas_validas=len(validas), filas_con_error=con_error, errores=errores
                ).model_dump(),
            )
        await _validar_limite_plan(session, gym_id, len(validas))

        ahora = datetime.now(timezone.utc)
        deportistas: List[Dict[str, Any]] = []
        membresias: List[Dict[str, Any]] = []
        for f in validas:
            dep_id = uuid4()
            deportistas.append({
                "id": dep_id,
                "gym_id": gym_id,
                "tipo_documento": f.tipo_documento,
                "documento": f.documento,
                "nombre": f.nombre,
                "correo": f.correo,
                "telefono": f.telefono,
                "fecha_nacimiento": f.fecha_nacimiento,
                "consentimiento_fecha": ahora,
                "consentimiento_version": CONSENTIMIENTO_VERSION_MIGRACION,
                "staff_id": staff_id,
                "ahora": ahora,
            })
            if f.plan is not None:
                duracion = f.plan["duracion_dias"]
                membresias.append({
                    "gym_id": gym_id,
                    "deportista_id": dep_id,
                    "plan_id": f.plan["id"],
                    # Caja recalcula el vencimiento como inicio + duración vendida: se guarda el
                    # inicio hacia atrás para que una renovación posterior parta del vencimiento real.
                    "fecha_inicio": f.fecha_vencimiento - timedelta(days=duracion),
                    "fecha_vencimiento": f.fecha_vencimiento,
                    "precio": f.plan["precio"],
                    "duracion": duracion,
                    "ahora": ahora,
                })

        try:
            async with session.begin_nested():
                await session.execute(
                    text("""
                        INSERT INTO platform.deportistas (
                            id, gimnasio_id, tipo_documento, documento, nombre, correo, telefono,
                            fecha_nacimiento, consentimiento_1581, consentimiento_fecha,
                            consentimiento_version, consentimiento_registrado_por,
                            activo, version, created_at, updated_at
                        ) VALUES (
                            :id, :gym_id, :tipo_documento, :documento, :nombre, :correo, :telefono,
                            :fecha_nacimiento, true, :consentimiento_fecha,
                            :consentimiento_version, :staff_id,
                            true, 1, :ahora, :ahora
                        )
                    """),
                    deportistas,
                )
                if membresias:
                    await session.execute(
                        text("""
                            INSERT INTO platform.membresias (
                                gimnasio_id, deportista_id, plan_id, fecha_inicio, fecha_vencimiento,
                                plan_precio_vendido, plan_duracion_dias_vendida, origen, created_at, updated_at
                            ) VALUES (
                                :gym_id, :deportista_id, :plan_id, :fecha_inicio, :fecha_vencimiento,
                                :precio, :duracion, 'migracion', :ahora, :ahora
                            )
                        """),
                        membresias,
                    )
        except IntegrityError as exc:
            if "ux_dep_documento_gym" in str(getattr(exc, "orig", exc)):
                # Otro usuario registró uno de estos documentos mientras se cargaba el archivo.
                raise _error(
                    status.HTTP_409_CONFLICT,
                    "DOCUMENTO_YA_REGISTRADO",
                    "Uno de los documentos se registró en el gimnasio mientras se cargaba el archivo. "
                    "Vuelve a simular la carga.",
                )
            raise

        await AuditService.registrar(
            session=session,
            gimnasio_id=gym_id,
            actor_id=staff_id,
            actor_nombre=staff_nombre,
            accion="CARGA_INICIAL",
            entidad="deportistas",
            entidad_id=None,
            detalle={
                "filas": len(filas),
                "deportistas_creados": len(deportistas),
                "membresias_creadas": len(membresias),
                "consentimiento_declarado": True,
            },
        )

        return AplicacionCargaResponse(
            deportistas_creados=len(deportistas), membresias_creadas=len(membresias)
        )
